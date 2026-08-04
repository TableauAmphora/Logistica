# conexiones para el proyecto de inventarios.

import truststore
truststore.inject_into_ssl()  # usa el almacen de certificados de Windows (proxy corporativo intercepta TLS), igual que schannel en git

import psycopg2
import pandas as pd 
from sqlalchemy import create_engine
from typing import Literal
from hdbcli import dbapi
from dotenv import load_dotenv
import sys
import requests
from lxml import etree
# import xml.etree.ElementTree as ET
# import time
import os
from datetime import datetime, timedelta
# import msal 
from openpyxl import load_workbook
# from openpyxl.styles import Font, Border, Side
# from io import BytesIO
# import base64
from config import ruta_credenciales


# Cargar el archivo .env desde la carpeta config
load_dotenv(dotenv_path= ruta_credenciales)








literal_bbdd_fuentes = Literal["dwh", "sap", "wms"]
literal_bbdd_salidas = Literal["dwh"]
class NoDatabaseConfiguredError(Exception):
    pass


def get_conexion(bbdd : literal_bbdd_fuentes): 
    f""" esta es la documentacion {literal_bbdd_fuentes} """
    if bbdd == 'dwh': 
        conexion = psycopg2.connect(host=os.getenv("dwh_server"), database = os.getenv("dwh_name") , user= os.getenv("dwh_user"), password= os.getenv("dwh_password"), port= os.getenv("dwh_port"))
    elif bbdd == 'sap':
        conexion = dbapi.connect( address= os.getenv("sap_server"), port= os.getenv("sap_port"),user= os.getenv("sap_user"), password= os.getenv("sap_password"))
    elif bbdd == 'wms':
        conexion = psycopg2.connect(host=os.getenv("wms_server"), database = os.getenv("wms_name") , user= os.getenv("wms_user"), password= os.getenv("wms_password"), port= os.getenv("wms_port"))
    else: 
        raise NoDatabaseConfiguredError(f"No se configuró conexión a {bbdd}.")
    return conexion 

def get_engine(bbdd:literal_bbdd_salidas): 
    if bbdd == 'dwh':
        engine = create_engine(f'postgresql://{os.getenv("dwh_user")}:{os.getenv("dwh_password")}@{os.getenv("dwh_server")}:{os.getenv("dwh_port")}/{os.getenv("dwh_name")}') 
    else:
        raise NoDatabaseConfiguredError(f"No se configuró engine a {bbdd}.")
    return engine 


def get_consulta(bbdd : literal_bbdd_fuentes, query, params = ()): 
    conexion = get_conexion(bbdd)
    if params == ():
        df = pd.read_sql(query, conexion)
    else:
        df = pd.read_sql(query, conexion, params = params)
    # cursor = conexion.cursor()
    # cursor.execute(query, params)
    # data = cursor.fetchall()
    # name_col = [cursor.description[j][0] for j in range(len(cursor.description))]
    # if len(data) == 0:
    #     print(f"no se han obtenido datos")
    #     df = pd.DataFrame(columns = name_col)
    #     return df
    # df = pd.DataFrame(data, columns = name_col) 
    
    return df 



def ejecutar_query(bbdd : literal_bbdd_salidas, query, params = ()):
    conexion = get_conexion(bbdd)
    cursor = conexion.cursor()
    cursor.execute(query, params)
    conexion.commit()
    cursor.close()
    # print("Se ejecuto la query con exito!") 


def cargar_log( flujo, fecha_inicio, fecha_fin, filas_procesadas = None, filas_nuevas = None, filas_actualizadas = None, estado = 'FAIL', mensaje = None, exit = True):
    log = {  "flujo_id" : [flujo], 
           "fecha_inicio" : [fecha_inicio], "fecha_fin" : [fecha_fin],
             "filas_procesadas" : [filas_procesadas], "filas_nuevas" : [filas_nuevas], 
             "filas_actualizadas" : [filas_actualizadas],
            "estado" : [estado], "mensaje" :  [mensaje]}
    log =  pd.DataFrame(log)
    engine = get_engine('dwh')
    log.to_sql('flujos_logs', engine, schema= 'control',  index=False, if_exists='append')
    
    sql_log = """
        SELECT id
        FROM control.flujos_logs
        WHERE flujo_id = %s
        ORDER BY id DESC
        LIMIT 1
        """
    # print(sql_log)
    log_id = pd.read_sql(
        sql_log,
        engine, params = (flujo,)
    ).iloc[0, 0]
    
    if exit == True:
        sys.exit(1) 
    return log_id 

def cargar_errores(df): 
    df.to_sql('flujos_errores', get_engine('dwh'), if_exists = 'append', index = False, schema = 'control')


# # #### CORREO ####  

# correo_tenant_id = os.getenv("correo_tenant_id")
# correo_client_id = os.getenv("correo_client_id")
# correo_client_secret = os.getenv("correo_client_secret")


# AUTHORITY = f'https://login.microsoftonline.com/{correo_tenant_id}' 
# user_default = 'fgonzalez@amphora.cl'
# id_carpeta_destino = 'AAMkADM1ZDY1MTY4LWNkMTUtNDI0Yi05ZjU1LWZhNGM2MzQwZGJlZAAuAAAAAAALK0hhc1L6ToKKLgS-m1SWAQAiMMDxrDLQTakDw9Gl0dXzAAExpS3LAAA=' ## Datos Sell out
# SCOPE = ['https://graph.microsoft.com/.default']




# def df_to_base64(df, sheet_name = 'Hoja1'): 
#     df_buffer = BytesIO()
#     with pd.ExcelWriter(df_buffer, engine='openpyxl') as writer:
#         df.to_excel(writer, sheet_name=sheet_name, index=False)
#     df_buffer.seek(0) 
#     wb = load_workbook(df_buffer)
#     ws = wb[sheet_name]  
#     ws.auto_filter.ref = ws.dimensions 
#     column_widths = {}
#     for row in ws.iter_rows():
#         for cell in row:
#             if cell.column_letter not in column_widths:
#                 column_widths[cell.column_letter] = len(str(cell.value)) 
#             else:
#                 column_widths[cell.column_letter] = max(column_widths[cell.column_letter], len(str(cell.value)))

#     for col, width in column_widths.items():
#         ws.column_dimensions[col].width = width + 4 
    
#     # Estilo normal (sin negrita, sin bordes)
#     normal_font = Font(bold=False)
#     no_border = Border(left=Side(style=None),    right=Side(style=None),    top=Side(style=None),    bottom=Side(style=None))

#     for cell in ws[1]:
#         cell.font = normal_font
#         cell.border = no_border

#     final_df_buffer = BytesIO()
#     wb.save(final_df_buffer)
#     final_df_buffer.seek(0)
#     df_base64 = base64.b64encode(final_df_buffer.read()).decode('utf-8')
#     return df_base64

# def get_token():
#     app = msal.ConfidentialClientApplication(correo_client_id, authority=AUTHORITY, client_credential=correo_client_secret )
#     result = app.acquire_token_for_client(scopes=SCOPE) 

#     if 'access_token' not in result:
#         print("Error al obtener token:", result.get("error_description"))
#         exit(1)

#     access_token = result['access_token']
#     return access_token


# def get_correos(user = user_default, token = get_token(), carpeta = 'inbox', fecha_desde = (datetime.now() - timedelta(days=5)).strftime('%Y-%m-%dT%H:%M:%SZ'), fecha_hasta = datetime.now().strftime('%Y-%m-%dT%H:%M:%SZ'), top = 500 ):
#     """ fecha_desde = (datetime.now() - timedelta(days=5)).strftime('%Y-%m-%dT%H:%M:%SZ') """
#     # Define headers para Graph
#     headers = {'Authorization': f'Bearer {token}', 'Content-Type': 'application/json'}
#     url = f'https://graph.microsoft.com/v1.0/users/{user}/mailFolders/{carpeta}/messages'


#     filter_query = (f"receivedDateTime ge {fecha_desde} and receivedDateTime le {fecha_hasta}   and contains(subject, 'Ventas')")
#     filter_query = (f"receivedDateTime ge {fecha_desde} and receivedDateTime le {fecha_hasta} ")

#     params = { '$filter': filter_query,
#             '$select': 'subject,receivedDateTime,bodyPreview,from',
#             '$orderby': 'receivedDateTime desc',
#             '$top': top  
#             }

#     response = requests.get(url, headers=headers, params=params)


    
#     if response.status_code != 200:
#         return f"""Error en la petición: status_code = {response.status_code} \n\n response.text = {response.text}"""
#         exit(1)
#     # print(f"""Error en la petición: status_code = {response.status_code} \n\n response.text = {response.text}""")
#     return response.json()

# def mover_correo(token, message_id, id_carpeta_destino, user_id, nombre_carpeta = 'Datos Sell Out'):
#     url = f"https://graph.microsoft.com/v1.0/users/{user_id}/messages/{message_id}/move"
#     headers = { "Authorization": f"Bearer {token}", "Content-Type": "application/json"}
#     body = { "destinationId": id_carpeta_destino}

#     response = requests.post(url, headers=headers, json=body)
#     if response.status_code == 201:
#         m = f"""mensaje movido a la carpeta {nombre_carpeta} exitosamente."""
#     else:
#         m = f"""❌❌❌ el mensaje no ha sido movido. Status code: {response.status_code} error : {response.json()}"""
#     return m


# def enviar_correo(mensaje, asunto, usuario = user_default, token = get_token() , destinatarios = [user_default], en_copia = [], en_copia_oculta = [], archivos = {}, tipo_mensaje = 'Text'):
#     to_recipients = [{"emailAddress": {"address": email}} for email in destinatarios]
#     cc_recipients = [{"emailAddress": {"address": email}} for email in en_copia]
#     bcc_recipients = [{"emailAddress": {"address": email}} for email in en_copia_oculta]
#     archivos_correo = [{
#                         "@odata.type": "#microsoft.graph.fileAttachment",
#                         "name": nombre_archivo,
#                         "contentType": "application/octet-stream",
#                         "contentBytes": archivo_base64
#                         }
#                         for nombre_archivo, archivo_base64 in archivos.items()
#                       ]

#     message = {  "message": { "subject": asunto,
#                             "body": { "contentType": tipo_mensaje,
#                                         "content": mensaje},
#                                 "toRecipients": to_recipients,
#                                 "ccRecipients": cc_recipients,
#                                 "bccRecipients": bcc_recipients,  
#                                 "attachments": archivos_correo    
#                             }
#                }
#     headers = { "Authorization": f"Bearer {token}", "Content-Type": "application/json"}
#     endpoint = f'https://graph.microsoft.com/v1.0/users/{usuario}/sendMail'
#     response = requests.post(endpoint, headers=headers, json=message)

#     return response.status_code

# def get_adjuntos(correo_id, token, user_id):
#     estado = 'OK'
#     attach_url = f"https://graph.microsoft.com/v1.0/users/{user_id}/messages/{correo_id}/attachments"
#     headers = {"Authorization": f"Bearer {token}"}
#     response = requests.get(attach_url, headers=headers)
#     if response.status_code == 200:
#         attachments = response.json().get('value', [])
#         return attachments
#     mensaje = f"Problemas al obtener adjuntos. \n status = {response.status_code} \ndetalle: {response.text}" 
#     raise Exception(mensaje)  
    
# def adjunto_to_df():
#     return 0 






    






# ### TABLEAU ### 


def get_token_tableau():
    version_api = os.getenv("tableau_version_api")
    server = os.getenv("tableau_server")
    token_name = os.getenv("tableau_token_name")
    token_secret = os.getenv("tableau_token_secret")
    site = os.getenv("tableau_site")
    url = f"{server}/api/{version_api}/auth/signin"
    payload = {
        "credentials": {
            "personalAccessTokenName": token_name,
            "personalAccessTokenSecret": token_secret,
            "site": {"contentUrl": site}
        }}

    headers = {"Content-Type": "application/json"}
    response = requests.post(url, json=payload, headers=headers)
    xml_response = response.text
    root = etree.fromstring(xml_response.encode("utf-8"))

    # Namespace
    ns = {"t": "http://tableau.com/api"}
    token = root.xpath("//t:credentials/@token", namespaces=ns)[0]
    site_id = root.xpath("//t:site/@id", namespaces=ns)[0]
    return token, site_id

def actualizar_libro(wb_id, token, site_id):
    version_api = os.getenv("tableau_version_api")
    server = os.getenv("tableau_server")
    url = f"{server}/api/{version_api}/sites/{site_id}/workbooks/{wb_id}/refresh" 

    headers = {  "X-Tableau-Auth": token,   "Content-Type": "application/xml"}
    body = "<tsRequest/>"
    response = requests.post(url, headers=headers, data = body)
    response.raise_for_status()
    return response

# def monitorear_trabajo(job_id, token, site_id, time_inicio, time_rep): 
#     mensaje = ''
#     version_api = os.getenv("tableau_version_api")
#     server = os.getenv("tableau_server")
#     status_url = f"{server}/api/{version_api}/sites/{site_id}/jobs/{job_id}"
#     headers = {"X-Tableau-Auth": token }

#     time.sleep(time_inicio)

#     while True:
#         r = requests.get(status_url, headers=headers)
#         root = ET.fromstring(r.text)
#         ns = {"t": "http://tableau.com/api"}
#         job = root.find(".//t:job", ns)
#         if job.attrib["progress"] == "100":
#             if job.attrib["finishCode"] == "0":
#                 return mensaje
#             else: 
#                 d = {"mode" : job.attrib["mode"], "type" : job.attrib["type"], "progress" : job.attrib["progress"], "createdAt" : job.attrib["createdAt"],
#                     "startedAt" : job.attrib["startedAt"],"completedAt" : job.attrib["completedAt"],"finishCode" : job.attrib["finishCode"],
#                     "note" : root.find(".//t:notes", ns).text }
#                 mensaje = mensaje + f"Problemas al actualizar libro de trabajo en Tableau. \n {d}"
#         else: 
#             pass
#         time.sleep(time_rep)













if __name__ == '__main__':
    pass



	




